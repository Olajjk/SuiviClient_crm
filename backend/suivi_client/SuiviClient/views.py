from rest_framework import generics, status
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated, AllowAny, BasePermission
from rest_framework.pagination import PageNumberPagination
from rest_framework_simplejwt.tokens import RefreshToken
from django.contrib.auth import get_user_model
from django.db.models import Q, Sum
from django.utils import timezone
from datetime import timedelta

from .models import Client, Interaction, RendezVous, MessageFidelisation, Campagne, Prestation, HistoriqueClient
from .serializers import (
    UserSerializer, RegisterSerializer, ClientListSerializer, ClientDetailSerializer,
    InteractionSerializer, RendezVousSerializer, PrestationSerializer,
    MessageFidelisationSerializer, CampagneSerializer, HistoriqueClientSerializer, ProfilSerializer
)

User = get_user_model()


# ── Permissions ───────────────────────────────────────────────

class IsAdmin(BasePermission):
    def has_permission(self, request, view):
        return request.user.is_authenticated and request.user.role == 'admin'

class IsAdminOrReadOnly(BasePermission):
    def has_permission(self, request, view):
        if not request.user.is_authenticated:
            return False
        if request.method in ('GET', 'HEAD', 'OPTIONS'):
            return True
        return request.user.role == 'admin'


# ── Pagination ────────────────────────────────────────────────

class StandardPagination(PageNumberPagination):
    page_size            = 20
    page_size_query_param = 'page_size'
    max_page_size        = 100


# ── Auth ──────────────────────────────────────────────────────

class RegisterView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        serializer = RegisterSerializer(data=request.data)
        if serializer.is_valid():
            user    = serializer.save()
            refresh = RefreshToken.for_user(user)
            return Response({
                'access':  str(refresh.access_token),
                'refresh': str(refresh),
                'user':    UserSerializer(user).data
            }, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class LoginView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        email    = request.data.get('email', '').lower().strip()
        password = request.data.get('password', '')
        try:
            user = User.objects.get(email=email)
        except User.DoesNotExist:
            return Response({'error': 'Email ou mot de passe incorrect'}, status=401)
        if not user.check_password(password):
            return Response({'error': 'Email ou mot de passe incorrect'}, status=401)
        if not user.is_active:
            return Response({'error': 'Compte désactivé'}, status=401)
        refresh = RefreshToken.for_user(user)
        return Response({
            'access':  str(refresh.access_token),
            'refresh': str(refresh),
            'user':    UserSerializer(user).data
        })


class LogoutView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        try:
            token = RefreshToken(request.data.get('refresh'))
            token.blacklist()
            return Response({'message': 'Déconnecté'})
        except Exception:
            return Response({'error': 'Token invalide'}, status=400)


class MeView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return Response(UserSerializer(request.user).data)


class MotDePasseOublieView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        from django.contrib.auth.tokens import default_token_generator
        from django.utils.encoding import force_bytes
        from django.utils.http import urlsafe_base64_encode
        from .notifications import envoyer_email

        email = request.data.get('email', '').lower().strip()
        try:
            user = User.objects.get(email=email)
        except User.DoesNotExist:
            # Même réponse pour éviter l'énumération d'emails
            return Response({'message': 'Si cet email existe, un lien a été envoyé.'})

        uid   = urlsafe_base64_encode(force_bytes(user.pk))
        token = default_token_generator.make_token(user)
        lien  = f"http://localhost:5173/reset-password/{uid}/{token}"

        html = f"""
        <div style="font-family:Arial;max-width:600px;margin:0 auto">
            <h2 style="color:#2563eb">Réinitialisation de mot de passe</h2>
            <p>Bonjour {user.first_name},</p>
            <p>Cliquez sur le lien ci-dessous pour réinitialiser votre mot de passe :</p>
            <a href="{lien}" style="display:inline-block;background:#2563eb;color:white;padding:12px 24px;border-radius:8px;text-decoration:none;margin:16px 0">
                Réinitialiser mon mot de passe
            </a>
            <p style="color:#6b7280;font-size:13px">Ce lien expire dans 24h. Si vous n'avez pas demandé ceci, ignorez cet email.</p>
        </div>
        """
        envoyer_email(user.email, 'Réinitialisation de votre mot de passe', html)
        return Response({'message': 'Si cet email existe, un lien a été envoyé.'})


class ResetMotDePasseView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        from django.contrib.auth.tokens import default_token_generator
        from django.utils.encoding import force_str
        from django.utils.http import urlsafe_base64_decode

        uid      = request.data.get('uid')
        token    = request.data.get('token')
        password = request.data.get('password')

        if not all([uid, token, password]):
            return Response({'error': 'Données manquantes'}, status=400)
        if len(password) < 8:
            return Response({'error': 'Mot de passe trop court (min 8 caractères)'}, status=400)

        try:
            pk   = force_str(urlsafe_base64_decode(uid))
            user = User.objects.get(pk=pk)
        except Exception:
            return Response({'error': 'Lien invalide'}, status=400)

        if not default_token_generator.check_token(user, token):
            return Response({'error': 'Lien expiré ou invalide'}, status=400)

        user.set_password(password)
        user.save()
        return Response({'message': 'Mot de passe mis à jour avec succès'})


# ── Dashboard ─────────────────────────────────────────────────

class DashboardView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        maintenant = timezone.now()
        debut_mois = maintenant.replace(day=1, hour=0, minute=0, second=0)
        user       = request.user
        is_admin   = user.role == 'admin'

        clients_qs     = Client.objects.all() if is_admin else Client.objects.filter(responsable=user)
        interactions_qs = Interaction.objects.all() if is_admin else Interaction.objects.filter(responsable=user)
        rdv_qs         = RendezVous.objects.all() if is_admin else RendezVous.objects.filter(responsable=user)

        prochains_rdv = rdv_qs.filter(
            date_debut__gte=maintenant,
            statut__in=['planifie', 'confirme']
        ).select_related('client')[:5]

        dernieres_interactions = interactions_qs.select_related('client')[:5]

        return Response({
            'stats': {
                'clients': {
                    'total':    clients_qs.count(),
                    'prospects': clients_qs.filter(statut='prospect').count(),
                    'actifs':   clients_qs.filter(statut='actif').count(),
                    'fideles':  clients_qs.filter(statut='fidele').count(),
                },
                'interactions': {
                    'ce_mois': interactions_qs.filter(created_at__gte=debut_mois).count(),
                    'total':   interactions_qs.count(),
                },
                'rdv': {
                    'a_venir': rdv_qs.filter(date_debut__gte=maintenant, statut__in=['planifie', 'confirme']).count(),
                    'total':   rdv_qs.count(),
                },
                'prestations': {
                    'en_cours': (Prestation.objects.all() if is_admin else Prestation.objects.filter(Q(responsable=user) | Q(client__responsable=user))).filter(statut='en_cours').count(),
                }
            },
            'prochains_rdv':          RendezVousSerializer(prochains_rdv, many=True).data,
            'dernieres_interactions': InteractionSerializer(dernieres_interactions, many=True).data,
            'prestations_en_cours':   PrestationSerializer(
                (Prestation.objects.all() if is_admin else Prestation.objects.filter(Q(responsable=user) | Q(client__responsable=user))).filter(statut='en_cours').select_related('client')[:5],
                many=True
            ).data,
        })


# ── Recherche globale ─────────────────────────────────────────

class RechercheGlobaleView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        q      = request.query_params.get('q', '').strip()
        user   = request.user
        is_admin = user.role == 'admin'

        if len(q) < 2:
            return Response({'clients': [], 'interactions': [], 'rdv': []})

        clients_qs     = Client.objects.all() if is_admin else Client.objects.filter(responsable=user)
        interactions_qs = Interaction.objects.all() if is_admin else Interaction.objects.filter(responsable=user)
        rdv_qs         = RendezVous.objects.all() if is_admin else RendezVous.objects.filter(responsable=user)

        clients = clients_qs.filter(
            Q(nom__icontains=q) | Q(prenom__icontains=q) |
            Q(email__icontains=q) | Q(code__icontains=q) | Q(telephone__icontains=q)
        )[:8]

        interactions = interactions_qs.filter(
            Q(sujet__icontains=q) | Q(description__icontains=q) | Q(client__nom__icontains=q)
        ).select_related('client')[:8]

        rdvs = rdv_qs.filter(
            Q(titre__icontains=q) | Q(lieu__icontains=q) | Q(client__nom__icontains=q)
        ).select_related('client')[:8]

        return Response({
            'clients':      ClientListSerializer(clients, many=True).data,
            'interactions': InteractionSerializer(interactions, many=True).data,
            'rdv':          RendezVousSerializer(rdvs, many=True).data,
        })


# ── Notifications in-app ──────────────────────────────────────

class NotificationsView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        user       = request.user
        maintenant = timezone.now()
        fin_auj    = maintenant.replace(hour=23, minute=59, second=59)

        rdv_qs = RendezVous.objects.filter(responsable=user) if user.role != 'admin' else RendezVous.objects.all()
        int_qs = Interaction.objects.filter(responsable=user) if user.role != 'admin' else Interaction.objects.all()

        rdvs_auj = rdv_qs.filter(
            date_debut__gte=maintenant,
            date_debut__lte=fin_auj,
            statut__in=['planifie', 'confirme']
        ).select_related('client')

        interactions_auj = int_qs.filter(
            date__gte=maintenant,
            date__lte=fin_auj,
            statut='planifie'
        ).select_related('client')

        notifs = []
        for rdv in rdvs_auj:
            notifs.append({
                'type':   'rdv',
                'titre':  rdv.titre,
                'client': rdv.client.nom,
                'heure':  rdv.date_debut.strftime('%H:%M'),
                'id':     rdv.id,
            })
        for inter in interactions_auj:
            notifs.append({
                'type':   'interaction',
                'titre':  inter.sujet,
                'client': inter.client.nom,
                'heure':  inter.date.strftime('%H:%M'),
                'id':     inter.id,
            })

        notifs.sort(key=lambda x: x['heure'])
        return Response({'notifications': notifs, 'total': len(notifs)})


# ── Clients ───────────────────────────────────────────────────

class ClientListView(generics.ListCreateAPIView):
    permission_classes = [IsAuthenticated]
    pagination_class   = StandardPagination

    def get_serializer_class(self):
        return ClientDetailSerializer if self.request.method == 'POST' else ClientListSerializer

    def get_queryset(self):
        user     = self.request.user
        is_admin = user.role == 'admin'

        if is_admin:
            qs = Client.objects.all()
        else:
            qs = Client.objects.filter(
                Q(responsable=user) | Q(responsable__isnull=True)
            )

        search = self.request.query_params.get('search')
        statut = self.request.query_params.get('statut')

        if search:
            qs = qs.filter(
                Q(nom__icontains=search) | Q(prenom__icontains=search) |
                Q(email__icontains=search) | Q(code__icontains=search)
            )
        if statut and statut != 'tous':
            qs = qs.filter(statut=statut)

        return qs

    def perform_create(self, serializer):
        if self.request.user.role != 'admin':
            serializer.save(responsable=self.request.user)
        else:
            serializer.save()


class ClientDetailView(generics.RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class   = ClientDetailSerializer

    def get_queryset(self):
        user = self.request.user
        if user.role == 'admin':
            return Client.objects.all()
        return Client.objects.filter(Q(responsable=user) | Q(responsable__isnull=True))

    def destroy(self, request, *args, **kwargs):
        return super().destroy(request, *args, **kwargs)


class ClientHistoriqueView(generics.ListAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class   = HistoriqueClientSerializer

    def get_queryset(self):
        return HistoriqueClient.objects.filter(client_id=self.kwargs['pk'])


# ── Interactions ──────────────────────────────────────────────

class InteractionListView(generics.ListCreateAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class   = InteractionSerializer
    pagination_class   = StandardPagination

    def get_queryset(self):
        user     = self.request.user
        is_admin = user.role == 'admin'
        qs = Interaction.objects.select_related('client', 'responsable')

        if not is_admin:
            qs = qs.filter(Q(responsable=user) | Q(client__responsable=user))

        client_id = self.request.query_params.get('client')
        search    = self.request.query_params.get('search')
        type_i    = self.request.query_params.get('type')

        if client_id: qs = qs.filter(client_id=client_id)
        if search:    qs = qs.filter(Q(sujet__icontains=search) | Q(client__nom__icontains=search))
        if type_i:    qs = qs.filter(type_interaction=type_i)

        return qs

    def perform_create(self, serializer):
        interaction = serializer.save(responsable=self.request.user)
        HistoriqueClient.objects.create(
            client=interaction.client,
            utilisateur=self.request.user,
            action=f"Interaction ajoutée : {interaction.sujet}"
        )


class InteractionDetailView(generics.RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class   = InteractionSerializer
    queryset           = Interaction.objects.all()


# ── Rendez-vous ───────────────────────────────────────────────

class RendezVousListView(generics.ListCreateAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class   = RendezVousSerializer
    pagination_class   = StandardPagination

    def get_queryset(self):
        user     = self.request.user
        is_admin = user.role == 'admin'
        qs = RendezVous.objects.select_related('client', 'responsable')

        if not is_admin:
            qs = qs.filter(Q(responsable=user) | Q(client__responsable=user))

        client_id = self.request.query_params.get('client')
        statut    = self.request.query_params.get('statut')
        a_venir   = self.request.query_params.get('a_venir')

        if client_id:        qs = qs.filter(client_id=client_id)
        if statut:           qs = qs.filter(statut=statut)
        if a_venir == 'true': qs = qs.filter(date_debut__gte=timezone.now(), statut__in=['planifie', 'confirme'])

        return qs

    def perform_create(self, serializer):
        rdv = serializer.save(responsable=self.request.user)
        HistoriqueClient.objects.create(
            client=rdv.client,
            utilisateur=self.request.user,
            action=f"RDV créé : {rdv.titre} le {rdv.date_debut.strftime('%d/%m/%Y')}"
        )
        # Email de confirmation au client (RDV futur uniquement)
        if rdv.client.email and rdv.date_debut > timezone.now():
            try:
                from .notifications import envoyer_confirmation_rdv
                envoyer_confirmation_rdv(rdv)
            except Exception:
                pass


class RendezVousDetailView(generics.RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class   = RendezVousSerializer
    queryset           = RendezVous.objects.all()

    def perform_update(self, serializer):
        ancien_statut = self.get_object().statut
        rdv = serializer.save()
        HistoriqueClient.objects.create(
            client=rdv.client,
            utilisateur=self.request.user,
            action=f"RDV modifié : {rdv.titre} — {ancien_statut} → {rdv.statut}"
        )
        # Email si statut changé ET RDV encore futur
        if rdv.statut != ancien_statut and rdv.client.email and rdv.date_debut > timezone.now():
            try:
                from .notifications import envoyer_notif_modif_rdv
                envoyer_notif_modif_rdv(rdv, ancien_statut)
            except Exception:
                pass

    def perform_destroy(self, instance):
        # Email d'annulation seulement si RDV futur avec statut actif
        rdv_futur    = instance.date_debut > timezone.now()
        statut_actif = instance.statut in ('planifie', 'confirme', 'reporte')
        if rdv_futur and statut_actif and instance.client.email:
            try:
                from .notifications import envoyer_notif_annulation_rdv
                envoyer_notif_annulation_rdv(instance)
            except Exception:
                pass
        HistoriqueClient.objects.create(
            client=instance.client,
            utilisateur=self.request.user,
            action=f"RDV supprimé : {instance.titre} du {instance.date_debut.strftime('%d/%m/%Y')}"
        )
        instance.delete()


class NotifierRdvView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            rdv = RendezVous.objects.get(pk=pk)
        except RendezVous.DoesNotExist:
            return Response({'error': 'RDV introuvable'}, status=404)

        from .notifications import envoyer_rappel_rdv
        succes = envoyer_rappel_rdv(rdv)
        if succes:
            rdv.notification_envoyee = True
            rdv.save()
            return Response({'message': 'Rappel envoyé'})
        return Response({'error': "Échec de l'envoi"}, status=500)



class ConfirmerRdvView(APIView):
    """
    Endpoint public (sans auth) — le client clique sur le lien dans son email.
    Passe le RDV à 'confirme' et retourne une page HTML de succès.
    """
    permission_classes = []  # accès public, pas besoin d'être connecté

    def get(self, request, token):
        from django.http import HttpResponse
        from django.utils import timezone

        try:
            rdv = RendezVous.objects.select_related('client', 'responsable').get(
                token_confirmation=token
            )
        except RendezVous.DoesNotExist:
            return HttpResponse(self._page_html(
                "❌ Lien invalide",
                "Ce lien de confirmation est invalide ou a expiré.",
                "#ef4444"
            ), content_type='text/html; charset=utf-8')

        if rdv.statut == 'annule':
            return HttpResponse(self._page_html(
                "⚠️ RDV annulé",
                "Ce rendez-vous a été annulé. Veuillez contacter votre responsable.",
                "#f59e0b"
            ), content_type='text/html; charset=utf-8')

        deja_confirme = rdv.statut == 'confirme'

        if not deja_confirme:
            rdv.statut      = 'confirme'
            rdv.confirme_le = timezone.now()
            rdv.save(update_fields=['statut', 'confirme_le'])

            # Notifier le responsable que le client a confirmé
            HistoriqueClient.objects.create(
                client=rdv.client,
                utilisateur=None,
                action=f"RDV confirmé par le client : {rdv.titre}"
            )
            # Email au responsable
            if rdv.responsable and rdv.responsable.email:
                from .notifications import envoyer_email, base_email
                date = rdv.date_debut.strftime('%d/%m/%Y à %Hh%M')
                html = base_email("RDV confirmé ✅", f"""
                  <h2 style="margin:0 0 8px;font-size:20px;font-weight:700;color:#18181b;">Confirmation reçue</h2>
                  <p style="margin:0 0 20px;font-size:15px;color:#71717a;">
                    <strong style="color:#18181b;">{rdv.client.nom} {rdv.client.prenom or ''}</strong>
                    a confirmé sa présence au rendez-vous.
                  </p>
                  <table width="100%" cellpadding="0" cellspacing="0"
                         style="border:1px solid #e4e4e7;border-radius:8px;overflow:hidden;">
                    <tr><td style="padding:12px 16px;background:#f9fafb;font-size:13px;
                                   font-weight:600;color:#71717a;border-bottom:1px solid #e4e4e7;">Objet</td>
                        <td style="padding:12px 16px;font-size:14px;color:#18181b;">{rdv.titre}</td></tr>
                    <tr><td style="padding:12px 16px;background:#f9fafb;font-size:13px;
                                   font-weight:600;color:#71717a;">Date</td>
                        <td style="padding:12px 16px;font-size:14px;color:#18181b;">{date}</td></tr>
                  </table>
                """)
                envoyer_email(rdv.responsable.email, f"✅ RDV confirmé par {rdv.client.nom}", html)

        date_str = rdv.date_debut.strftime('%d/%m/%Y à %Hh%M')
        message  = "Votre présence a bien été enregistrée." if not deja_confirme else "Vous avez déjà confirmé ce rendez-vous."

        return HttpResponse(self._page_html(
            "✅ Rendez-vous confirmé",
            f"{message}<br><br>"
            f"<strong>{rdv.titre}</strong><br>"
            f"📅 {date_str}"
            + (f"<br>📍 {rdv.lieu}" if rdv.lieu else ""),
            "#10b981"
        ), content_type='text/html; charset=utf-8')

    def _page_html(self, titre, message, couleur):
        return f"""<!DOCTYPE html>
<html lang="fr">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>{titre}</title>
  <style>
    * {{ margin: 0; padding: 0; box-sizing: border-box; }}
    body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', sans-serif;
            background: #f4f4f5; min-height: 100vh;
            display: flex; align-items: center; justify-content: center; padding: 20px; }}
    .card {{ background: white; border-radius: 16px; padding: 48px 40px;
             max-width: 480px; width: 100%; text-align: center;
             box-shadow: 0 4px 24px rgba(0,0,0,0.08); }}
    .icon {{ font-size: 56px; margin-bottom: 20px; }}
    h1 {{ font-size: 24px; font-weight: 700; color: #18181b; margin-bottom: 16px; }}
    p {{ font-size: 15px; color: #71717a; line-height: 1.7; }}
    .badge {{ display: inline-block; margin-top: 24px; padding: 8px 20px;
              background: {couleur}15; color: {couleur};
              border-radius: 99px; font-size: 13px; font-weight: 600; }}
  </style>
</head>
<body>
  <div class="card">
    <div class="icon">{'✅' if '✅' in titre else '❌' if '❌' in titre else '⚠️'}</div>
    <h1>{titre.replace('✅ ', '').replace('❌ ', '').replace('⚠️ ', '')}</h1>
    <p>{message}</p>
    <span class="badge">SMART PILOTE</span>
  </div>
</body>
</html>"""


# ── Prestations ───────────────────────────────────────────────

class PrestationListView(generics.ListCreateAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class   = PrestationSerializer
    pagination_class   = StandardPagination

    def get_queryset(self):
        user     = self.request.user
        is_admin = user.role == 'admin'
        qs = Prestation.objects.select_related('client', 'responsable')

        if not is_admin:
            qs = qs.filter(Q(responsable=user) | Q(client__responsable=user))

        client_id = self.request.query_params.get('client')
        statut    = self.request.query_params.get('statut')
        search    = self.request.query_params.get('search')

        if client_id: qs = qs.filter(client_id=client_id)
        if statut:    qs = qs.filter(statut=statut)
        if search:    qs = qs.filter(Q(nom__icontains=search) | Q(client__nom__icontains=search))

        return qs

    def perform_create(self, serializer):
        prestation = serializer.save(responsable=self.request.user)
        HistoriqueClient.objects.create(
            client=prestation.client,
            utilisateur=self.request.user,
            action=f"Prestation ajoutée : {prestation.nom} — {prestation.prix_final}€"
        )


class PrestationDetailView(generics.RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class   = PrestationSerializer
    queryset           = Prestation.objects.all()

    def perform_update(self, serializer):
        ancien_statut = self.get_object().statut
        prestation = serializer.save()
        if ancien_statut != 'termine' and prestation.statut == 'termine':
            HistoriqueClient.objects.create(
                client=prestation.client,
                utilisateur=self.request.user,
                action=f"Prestation terminée : {prestation.nom} — {prestation.prix_final}€"
            )
            from .notifications import envoyer_notif_prestation_terminee
            envoyer_notif_prestation_terminee(prestation)


# ── Fidélisation ──────────────────────────────────────────────

class MessageFidelisationListView(generics.ListCreateAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class   = MessageFidelisationSerializer

    def get_queryset(self):
        qs        = MessageFidelisation.objects.select_related('client')
        client_id = self.request.query_params.get('client')
        canal     = self.request.query_params.get('canal')
        if client_id: qs = qs.filter(client_id=client_id)
        if canal:     qs = qs.filter(canal=canal)
        return qs

    def perform_create(self, serializer):
        from .notifications import envoyer_message_fidelisation
        message = serializer.save()
        envoyer_message_fidelisation(message)


class MessageFidelisationDetailView(generics.RetrieveDestroyAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class   = MessageFidelisationSerializer
    queryset           = MessageFidelisation.objects.all()


# ── Campagnes (admin uniquement pour créer/modifier/supprimer) ─

class CampagneListView(generics.ListCreateAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class   = CampagneSerializer
    queryset           = Campagne.objects.all()

    def perform_create(self, serializer):
        if self.request.user.role != 'admin':
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied("Seul un administrateur peut créer une campagne.")
        serializer.save()


class CampagneDetailView(generics.RetrieveUpdateDestroyAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class   = CampagneSerializer
    queryset           = Campagne.objects.all()

    def update(self, request, *args, **kwargs):
        if request.user.role != 'admin':
            return Response({'error': 'Accès refusé'}, status=403)
        return super().update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        if request.user.role != 'admin':
            return Response({'error': 'Accès refusé'}, status=403)
        return super().destroy(request, *args, **kwargs)


class ExecuterCampagneView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            campagne = Campagne.objects.get(pk=pk)
        except Campagne.DoesNotExist:
            return Response({'error': 'Campagne introuvable'}, status=404)
        from .tasks import executer_campagne
        executer_campagne.delay(campagne.id)
        return Response({'message': f'Campagne "{campagne.nom}" lancée'})


class ExecuterToutesCampagnesView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        if request.user.role != 'admin':
            return Response({'error': 'Accès refusé'}, status=403)
        campagnes = Campagne.objects.filter(statut='active')
        from .tasks import executer_campagne
        for c in campagnes:
            executer_campagne.delay(c.id)
        return Response({'message': f'{campagnes.count()} campagne(s) lancée(s)'})


# ── PDF ───────────────────────────────────────────────────────

class ClientPDFView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            client = Client.objects.get(pk=pk)
        except Client.DoesNotExist:
            return Response({'error': 'Client introuvable'}, status=404)
        from .pdf_generator import generer_fiche_client_pdf
        return generer_fiche_client_pdf(client)


# ── Export Excel ──────────────────────────────────────────────

class ExportClientsExcelView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        import openpyxl
        from django.http import HttpResponse

        user     = request.user
        is_admin = user.role == 'admin'
        clients  = Client.objects.all() if is_admin else Client.objects.filter(
            Q(responsable=user) | Q(responsable__isnull=True)
        )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Clients'

        from openpyxl.styles import Font, PatternFill, Alignment
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(fill_type='solid', fgColor='2563EB')

        headers = ['Code', 'Nom', 'Prénom', 'Email', 'Téléphone', 'Ville', 'Statut', 'Responsable', 'Créé le']
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font  = header_font
            cell.fill  = header_fill
            cell.alignment = Alignment(horizontal='center')

        for row, c in enumerate(clients, 2):
            ws.append([
                c.code,
                c.nom,
                c.prenom or '',
                c.email or '',
                c.telephone or '',
                c.ville or '',
                c.get_statut_display(),
                c.responsable.get_full_name() if c.responsable else '',
                c.date_creation.strftime('%d/%m/%Y'),
            ])

        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 18

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="clients.xlsx"'
        wb.save(response)
        return response


class ExportPrestationsExcelView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        import openpyxl
        from django.http import HttpResponse

        user        = request.user
        is_admin    = user.role == 'admin'
        prestations = Prestation.objects.all() if is_admin else Prestation.objects.filter(
            Q(responsable=user) | Q(client__responsable=user)
        )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Prestations'

        from openpyxl.styles import Font, PatternFill, Alignment
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(fill_type='solid', fgColor='059669')

        headers = ['Client', 'Prestation', 'Date', 'Prix', 'Réduction %', 'Prix final', 'Statut', 'Responsable']
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font  = header_font
            cell.fill  = header_fill
            cell.alignment = Alignment(horizontal='center')

        for row, p in enumerate(prestations, 2):
            ws.append([
                p.client.nom,
                p.nom,
                p.date_realisation.strftime('%d/%m/%Y'),
                float(p.prix),
                float(p.reduction),
                float(p.prix_final),
                p.get_statut_display(),
                p.responsable.get_full_name() if p.responsable else '',
            ])

        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 16

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="prestations.xlsx"'
        wb.save(response)
        return response


# ── Profil ────────────────────────────────────────────────────

from rest_framework.parsers import MultiPartParser, FormParser

class ProfilView(APIView):
    permission_classes = [IsAuthenticated]
    parser_classes     = [MultiPartParser, FormParser]

    def get(self, request):
        return Response(UserSerializer(request.user).data)

    def put(self, request):
        serializer = ProfilSerializer(request.user, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data)
        return Response(serializer.errors, status=400)


class AvatarView(APIView):
    permission_classes = [IsAuthenticated]
    parser_classes     = [MultiPartParser, FormParser]

    def post(self, request):
        avatar = request.FILES.get('avatar')
        if not avatar:
            return Response({'error': 'Aucun fichier envoyé'}, status=400)
        if not avatar.content_type.startswith('image/'):
            return Response({'error': 'Le fichier doit être une image'}, status=400)
        if avatar.size > 5 * 1024 * 1024:
            return Response({'error': 'Image trop grande (max 5MB)'}, status=400)

        user = request.user
        if user.avatar:
            import os
            if os.path.exists(user.avatar.path):
                os.remove(user.avatar.path)

        user.avatar = avatar
        user.save(update_fields=['avatar'])
        return Response({
            'message':    'Avatar mis à jour',
            'avatar_url': request.build_absolute_uri(user.avatar.url)
        })